"""F57 — attachment read-aloud (PHASE_10_PLUS_SPEC.md §13.4). Real PDF/
.docx fixtures are generated on the fly, not checked-in binary files.
"""

import pytest
from docx import Document

from backend.attachment_reader import extract_text
from backend.errors import SpokenError


def _build_minimal_pdf(text: str) -> bytes:
    """A real, minimal single-page PDF with one text run, built with
    correct byte offsets (pypdf requires a working `startxref`, unlike
    some lenient readers — a hand-typed approximate one isn't enough).
    """
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        None,  # filled in below, once stream_content is known
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    stream_content = f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode("latin-1")
    objects[3] = (
        b"<</Length " + str(len(stream_content)).encode() + b">>stream\n"
        + stream_content + b"\nendstream"
    )

    body = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(body))
        body += f"{i} 0 obj".encode() + obj + b"endobj\n"

    xref_offset = len(body)
    xref = b"xref\n0 " + str(len(objects) + 1).encode() + b"\n"
    xref += b"0000000000 65535 f \n"
    for off in offsets:
        xref += f"{off:010d} 00000 n \n".encode()

    trailer = (
        b"trailer\n<</Size " + str(len(objects) + 1).encode() + b"/Root 1 0 R>>\n"
        b"startxref\n" + str(xref_offset).encode() + b"\n%%EOF"
    )
    return body + xref + trailer


def test_extract_text_pdf(tmp_path):
    path = tmp_path / "test.pdf"
    path.write_bytes(_build_minimal_pdf("Hello PDF World"))
    text = extract_text(str(path))
    assert "Hello PDF World" in text


def test_extract_text_docx(tmp_path):
    path = tmp_path / "test.docx"
    document = Document()
    document.add_paragraph("Hello Word World.")
    document.add_paragraph("A second paragraph.")
    document.save(str(path))

    text = extract_text(str(path))
    assert "Hello Word World." in text
    assert "A second paragraph." in text


def test_extract_text_plain_txt(tmp_path):
    path = tmp_path / "test.txt"
    path.write_text("Just plain text, unchanged.", encoding="utf-8")
    assert extract_text(str(path)) == "Just plain text, unchanged."


def test_extract_text_unsupported_type_names_it(tmp_path):
    path = tmp_path / "test.xlsx"
    path.write_bytes(b"not a real spreadsheet, just testing the type name")
    with pytest.raises(SpokenError) as exc_info:
        extract_text(str(path))
    assert "a spreadsheet" in exc_info.value.message


def test_extract_text_unknown_type_falls_back_to_generic_name(tmp_path):
    path = tmp_path / "test.weirdext"
    path.write_bytes(b"whatever")
    with pytest.raises(SpokenError) as exc_info:
        extract_text(str(path))
    assert "that kind of file" in exc_info.value.message


# ---------------------------------------------------------------------------
# Route-level
# ---------------------------------------------------------------------------


def _attach_txt_file(client, session_id: str, text: str) -> None:
    import io

    client.post(
        "/api/mail/attachment",
        data={"session_id": session_id},
        files={"file": ("notes.txt", io.BytesIO(text.encode("utf-8")), "text/plain")},
    )


def test_read_attachment_short_text_reads_directly(client):
    session_id = "attach-read-short"
    client.post("/api/turn", json={"session_id": session_id, "transcript": "tell John Smith I will be late"})
    _attach_txt_file(client, session_id, "A short note about the meeting.")
    client.post("/api/turn", json={"session_id": session_id, "transcript": "attach the last document i mentioned"})

    resp = client.post("/api/turn", json={"session_id": session_id, "transcript": "read the attachment"})
    body = resp.json()
    assert body["speech"] == "A short note about the meeting."


def test_read_attachment_long_text_offers_summary_then_full(client):
    session_id = "attach-read-long"
    client.post("/api/turn", json={"session_id": session_id, "transcript": "tell John Smith I will be late"})
    long_text = " ".join(["word"] * 350)
    _attach_txt_file(client, session_id, long_text)
    client.post("/api/turn", json={"session_id": session_id, "transcript": "attach the last document i mentioned"})

    resp = client.post("/api/turn", json={"session_id": session_id, "transcript": "read the attachment"})
    body = resp.json()
    assert "about 350 words" in body["speech"]
    assert "read it in full" in body["speech"]

    full = client.post("/api/turn", json={"session_id": session_id, "transcript": "read it in full"})
    assert full.json()["speech"] == long_text


def test_read_attachment_refused_without_a_draft(client):
    resp = client.post(
        "/api/turn", json={"session_id": "attach-no-draft", "transcript": "read the attachment"}
    )
    assert "only available" in resp.json()["speech"]


def test_read_attachment_with_no_attachment_yet(client):
    session_id = "attach-none"
    client.post("/api/turn", json={"session_id": session_id, "transcript": "tell John Smith I will be late"})
    resp = client.post("/api/turn", json={"session_id": session_id, "transcript": "read the attachment"})
    body = resp.json()
    assert "haven't attached anything yet" in body["speech"]


def test_read_full_still_means_inbox_message_when_no_attachment_pending(client):
    session_id = "attach-readfull-inbox"
    client.post("/api/turn", json={"session_id": session_id, "transcript": "read my unread mail"})
    resp = client.post("/api/turn", json={"session_id": session_id, "transcript": "read it in full"})
    body = resp.json()
    assert body["ok"] is True
    assert body["phase"] == "reading_inbox"
